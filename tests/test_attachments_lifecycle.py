"""Behavior tests for the migrated file lifecycle (no live services required)."""
from io import BytesIO
from datetime import datetime, timedelta
import pytest
import pytest_asyncio
from fastapi import UploadFile
from sqlmodel import SQLModel
from sqlalchemy.ext.asyncio import create_async_engine
from sqlmodel.ext.asyncio.session import AsyncSession
from backend.attachments.models import FilePurpose
from backend.attachments.service.file_resource_service import FileResourceService
from backend.attachments.service.upload_session_service import UploadSessionService
from backend.attachments.service.content_extractor import extract_text_from_bytes, chunk_text
from backend.attachments.api.references import set_references, ReferenceSet
from backend.attachments.vendor.file.store.file_store_helper import get_file_store
from backend.attachments.support import FileStatus

@pytest_asyncio.fixture
async def session(tmp_path, monkeypatch):
    monkeypatch.setenv('AITUGE_ATTACHMENT_STORAGE_ROOT', str(tmp_path / 'files'))
    get_file_store.cache_clear()
    engine = create_async_engine('sqlite+aiosqlite:///' + str(tmp_path / 'db.sqlite'))
    async with engine.begin() as conn:
        await conn.run_sync(SQLModel.metadata.create_all)
    async with AsyncSession(engine, expire_on_commit=False) as session:
        yield session
    await engine.dispose()
    get_file_store.cache_clear()

def upload(raw, name='file.txt'):
    return UploadFile(filename=name, file=BytesIO(raw))

@pytest.mark.asyncio
async def test_dedup_retry_refs_gc(session):
    svc = FileResourceService(session)
    async def create():
        return await svc.create_from_upload(upload=upload(b'hello'), purpose=FilePurpose.CHAT_ATTACHMENT, tenant_id='t')
    a, new = await create()
    b, reused = await create()
    assert new and not reused and a.id == b.id
    await svc.mark_status(file_id=a.id, tenant_id='t', status=FileStatus.failed, failed_reason='offline')
    revived, new = await create()
    assert new and revived.id == a.id and revived.failed_reason is None
    a.expires_at = datetime.utcnow() - timedelta(days=1)
    session.add(a)
    await session.commit()
    for _ in range(2):
        await set_references(session, 't', 'message-1', ReferenceSet(revision=1, file_ids=[a.id]))
    await session.refresh(a)
    assert a.ref_count == 1 and not await svc.sweep_expired_candidates()
    with pytest.raises(ValueError):
        await svc.hard_delete(a.id, 't')
    await set_references(session, 't', 'message-1', ReferenceSet(revision=2, file_ids=[]))
    await set_references(session, 't', 'message-1', ReferenceSet(revision=1, file_ids=[a.id]))
    await session.refresh(a)
    assert a.ref_count == 0
    path = get_file_store().base_path + '/' + a.file_path
    assert await svc.hard_delete(a.id, 't')
    from pathlib import Path
    assert not Path(path).exists()

@pytest.mark.asyncio
async def test_multipart_resume_and_expiry(session):
    svc = UploadSessionService(session)
    row = await svc.create_session(tenant_id='t', file_name='merged.txt', purpose=FilePurpose.CHAT_ATTACHMENT)
    for part, raw in [(1,b'first'), (1,b'first'), (2,b'second')]:
        await svc.write_part(upload_id=row.id, tenant_id='t', part_number=part, upload=upload(raw))
    resumed = await UploadSessionService(session).get_session(row.id, 't')
    assert len(resumed.parts) == 2
    file, new = await svc.complete(upload_id=row.id, tenant_id='t', expected_part_count=2)
    assert new
    same, new = await svc.complete(upload_id=row.id, tenant_id='t', expected_part_count=2)
    assert not new and same.id == file.id
    data = await FileResourceService(session).read_bytes(file.id, 't')
    assert data.read() == b'firstsecond'
    row = await svc.create_session(tenant_id='t', file_name='expired.txt', purpose=FilePurpose.CHAT_ATTACHMENT)
    row.expires_at = datetime.utcnow() - timedelta(seconds=1)
    await session.commit()
    with pytest.raises(ValueError, match='过期'):
        await svc.write_part(upload_id=row.id, tenant_id='t', part_number=1, upload=upload(b'part'))

@pytest.mark.asyncio
async def test_long_text_and_same_names(session):
    svc = FileResourceService(session)
    ids = []
    for raw in (b'a', b'b'):
        row, _ = await svc.create_from_upload(upload=upload(raw), purpose=FilePurpose.CHAT_ATTACHMENT, tenant_id='t')
        ids.append(row.id)
        text = '前文' * 6000 + '验收金额 1250 元'
        await svc.write_text_content(file_id=row.id, tenant_id='t', content=text, extractor_version='test')
        await svc.replace_chunks(file_id=row.id, tenant_id='t', chunks=chunk_text(text))
    assert len(await svc.get_file_contents_map(ids, 't')) == 2
    tail = await svc.get_text_slice(file_id=ids[0], tenant_id='t', offset=12000, limit=100)
    assert '1250' in tail['content']
    hits = await svc.search_chunks(file_id=ids[0], tenant_id='t', query='1250')
    assert hits[0]['start_offset'] > 10000

@pytest.mark.parametrize('extension', ['.docx', '.pptx', '.xlsx'])
def test_real_office_formats(extension):
    buf = BytesIO()
    if extension == '.docx':
        from docx import Document
        doc = Document(); doc.add_paragraph('附件验收文档'); doc.save(buf)
    elif extension == '.pptx':
        from pptx import Presentation
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text = '附件验收文档'; prs.save(buf)
    else:
        import pandas as pd
        with pd.ExcelWriter(buf, engine='openpyxl') as writer:
            pd.DataFrame({'数额':[100,200]}).to_excel(writer, sheet_name='第一表', index=False)
            pd.DataFrame({'数额':[900]}).to_excel(writer, sheet_name='附件验收文档', index=False)
    text, truncated = extract_text_from_bytes(buf.getvalue(), extension, file_name='sample'+extension)
    assert '附件验收文档' in text and not truncated
    if extension == '.xlsx': assert '900' in text and '第一表' in text

def test_legacy_and_corrupt_office_fail():
    for ext in ['.doc', '.ppt', '.docx', '.pptx']:
        with pytest.raises(Exception):
            extract_text_from_bytes(b'corrupt', ext)

@pytest.mark.asyncio
async def test_default_attachment_bundle_reads_images_without_mounting_visual(session, monkeypatch):
    from contextlib import asynccontextmanager
    import json
    import backend.attachments.tools as attachment_tools
    import backend.attachments.tools.file_reader as reader

    @asynccontextmanager
    async def db_session():
        yield session

    monkeypatch.setattr(attachment_tools, 'create_db_session', db_session)
    monkeypatch.setattr(reader, 'create_db_session', db_session)
    svc = FileResourceService(session)
    row, _ = await svc.create_from_upload(upload=upload(b'image-fixture', 'receipt.jpg'),
                                         purpose=FilePurpose.CHAT_ATTACHMENT, tenant_id='t')
    await svc.write_text_content(file_id=row.id, tenant_id='t', content='识别金额 200 元', extractor_version='test')
    await svc.mark_status(file_id=row.id, tenant_id='t', status=FileStatus.succeeded)
    bundle, prompt, _ = await attachment_tools.create_attachment_bundle([{'file_id': row.id}], 't')
    try:
        assert {tool.metadata.name for tool in bundle.tools} == {'read-file', 'search-file-chunks'}
        assert 'multimodal-parser' not in prompt
        read = next(tool for tool in bundle.tools if tool.metadata.name == 'read-file')
        result = await read.acall(file_id=row.id)
        assert json.loads(result.content)['content'] == '识别金额 200 元'
        await svc.write_text_content(file_id=row.id, tenant_id='t', content='', extractor_version='test')
        result = await read.acall(file_id=row.id)
        assert '未提取到' in result.content and 'multimodal-parser' not in result.content
    finally:
        await bundle.cleanup()
    # The implementation remains available for explicit integration, without
    # contacting a visual provider during bundle construction.
    bundle, _, _ = await attachment_tools.create_attachment_bundle([{'file_id': row.id}], 't', include_visual=True)
    try:
        assert 'multimodal-parser' in {tool.metadata.name for tool in bundle.tools}
    finally:
        await bundle.cleanup()
