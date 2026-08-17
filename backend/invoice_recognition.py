"""Typed invoice recognition for the travel reimbursement workflow."""
from __future__ import annotations

import hashlib
import hmac
import io
import os
import re
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Annotated

import httpx
from docx import Document
from fastapi import APIRouter, File, Header, HTTPException, UploadFile
from pydantic import BaseModel, ConfigDict, Field
from pypdf import PdfReader

MAX_FILES = 20
MAX_FILE_BYTES = 20 * 1024 * 1024
MIN_NATIVE_TEXT_LENGTH = 40
_MONEY = re.compile(r"(?<![\d.])-?\d{1,9}(?:,\d{3})*(?:\.\d{2})(?!\d)")
_INVOICE_NO = re.compile(r"(?<!\d)(\d{20})(?!\d)")
_TAX_ID = re.compile(r"(?<![0-9A-Z])([0-9A-Z]{18})(?![0-9A-Z])")
_DATES = (
    re.compile(r"(20\d{2})年\s*(\d{1,2})月\s*(\d{1,2})日"),
    re.compile(r"(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})"),
)
_COMPANY_SUFFIXES = (
    "有限责任公司", "股份有限公司", "集团有限公司", "科技有限公司",
    "软件有限公司", "餐饮有限公司", "管理有限公司", "有限公司",
    "医院", "大学", "学院", "中心",
)


class InvoiceItem(BaseModel):
    model_config = ConfigDict(extra="forbid")
    invoice_id: str
    file_name: str
    expense_category: str
    recognition_method: str
    invoice_type: str
    invoice_number: str | None = None
    issue_date: date | None = None
    buyer_name: str | None = None
    buyer_tax_id: str | None = None
    seller_name: str | None = None
    seller_tax_id: str | None = None
    amount_excluding_tax: Decimal | None = None
    tax_amount: Decimal | None = None
    total_amount: Decimal | None = None
    confidence: float = Field(ge=0, le=1)
    warnings: list[str] = Field(default_factory=list)


class InvoiceRecognitionResponse(BaseModel):
    model_config = ConfigDict(extra="forbid")
    invoices: list[InvoiceItem]
    total_amount: Decimal
    total_amount_excluding_tax: Decimal
    total_tax_amount: Decimal
    file_count: int
    recognized_count: int


def _check_token(value: str) -> None:
    expected = (
        os.getenv("FRAMEWORK_INTERNAL_TOKEN", "").strip()
        or os.getenv("CONTRACT_INTERNAL_TOKEN", "").strip()
    )
    if not expected or not hmac.compare_digest(value, expected):
        raise HTTPException(status_code=401, detail="Invalid internal credential")


def _pdf_text(content: bytes) -> str:
    try:
        return "\n".join(
            (page.extract_text() or "") for page in PdfReader(io.BytesIO(content)).pages
        )
    except Exception:
        return ""


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    lines = [item.text for item in document.paragraphs if item.text]
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(lines)


async def _ocr_text(file_name: str, content: bytes) -> str:
    base_url = os.getenv("INVOICE_OCR_BASE_URL", "").strip().rstrip("/")
    if not base_url:
        raise HTTPException(status_code=422, detail="OCR service is unavailable.")
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(180.0)) as client:
            response = await client.post(
                f"{base_url}/v1/internal/ocr/convert",
                files={"file": (file_name, content, "application/pdf")},
            )
            response.raise_for_status()
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            return _docx_text(archive.read("contract.docx"))
    except (httpx.HTTPError, KeyError, zipfile.BadZipFile, ValueError) as exc:
        raise HTTPException(status_code=502, detail="Invoice OCR failed.") from exc


def _parse_date(text: str) -> date | None:
    for pattern in _DATES:
        match = pattern.search(text)
        if match:
            try:
                return date(*(int(value) for value in match.groups()))
            except ValueError:
                pass
    return None


def _company_names(text: str) -> list[str]:
    suffixes = "|".join(re.escape(value) for value in _COMPANY_SUFFIXES)
    pattern = re.compile(rf"([^\n\r：:票]{{2,70}}?(?:{suffixes}))")
    result: list[str] = []
    for line in text.splitlines():
        for match in pattern.finditer(line):
            value = re.sub(
                r"^(?:购买方|销售方|销方|购方|名称)\s*[：:]?\s*",
                "",
                match.group(1),
            ).strip(" ：:")
            if len(value) >= 4 and value not in result:
                result.append(value)
    return result


def _tax_ids(text: str) -> list[str]:
    result: list[str] = []
    for value in _TAX_ID.findall(text.upper()):
        if value not in result:
            result.append(value)
    return result


def _decimal(value: str) -> Decimal | None:
    try:
        return Decimal(value.replace(",", ""))
    except (InvalidOperation, ValueError):
        return None


def _amounts(text: str) -> tuple[Decimal | None, Decimal | None, Decimal | None]:
    values = [
        value for raw in _MONEY.findall(text)
        if (value := _decimal(raw)) is not None and value >= 0
    ]
    if not values:
        return None, None, None
    totals: list[Decimal] = []
    for pattern in (
        r"(?:价税合计|小写)[^\d]{0,30}(\d[\d,]*\.\d{2})",
        r"[¥￥]\s*(\d[\d,]*\.\d{2})",
    ):
        totals.extend(
            value for raw in re.findall(pattern, text)
            if (value := _decimal(raw)) is not None
        )
    total = totals[-1] if totals else values[-1]
    candidates = [value for value in values if value <= total]
    for index, amount in enumerate(candidates):
        for tax in candidates[index + 1:]:
            if amount > tax and amount + tax == total:
                return amount, tax, total
    return None, None, total


def _expense_category(file_name: str, text: str) -> str:
    searchable = re.sub(r"\s+", "", f"{file_name}\n{text}").lower()
    category_keywords = (
        ("ACCOMMODATION", ("住宿", "酒店", "宾馆", "旅馆", "客房", "房费")),
        ("MEAL", ("餐饮", "餐费", "饭店", "餐厅", "食品", "酒楼", "小吃", "用餐")),
        ("TRANSPORT", (
            "打车", "出租车", "出租汽车", "网约车", "滴滴", "高德打车", "曹操出行",
            "机票", "航空", "火车票", "铁路", "高铁", "客运", "交通费", "过路费", "燃油费",
        )),
    )
    for category, keywords in category_keywords:
        if any(keyword in searchable for keyword in keywords):
            return category
    return "OTHER"


def _invoice_type(text: str) -> str:
    compact = re.sub(r"\s+", "", text)
    if "增值税专用发票" in compact:
        return "VAT_SPECIAL"
    if "增值税普通发票" in compact or "电子发票" in compact:
        return "VAT_ORDINARY"
    return "UNKNOWN"


def parse_invoice(
    file_name: str,
    text: str,
    method: str = "PDF_TEXT",
    invoice_id: str | None = None,
) -> InvoiceItem:
    normalized = text.replace("\u3000", " ").replace("\xa0", " ")
    resolved_invoice_id = invoice_id or hashlib.sha256(
        f"{file_name}\n{normalized}".encode("utf-8")
    ).hexdigest()
    numbers = _INVOICE_NO.findall(normalized)
    names = _company_names(normalized)
    tax_ids = _tax_ids(normalized)
    amount, tax, total = _amounts(normalized)
    values = {
        "invoice_number": numbers[0] if numbers else None,
        "issue_date": _parse_date(normalized),
        "buyer_name": names[0] if names else None,
        "seller_name": names[1] if len(names) > 1 else None,
        "total_amount": total,
    }
    warnings = [
        f"MISSING_{key.upper()}" for key, value in values.items() if value is None
    ]
    return InvoiceItem(
        invoice_id=resolved_invoice_id,
        file_name=file_name,
        expense_category=_expense_category(file_name, normalized),
        recognition_method=method,
        invoice_type=_invoice_type(normalized),
        invoice_number=values["invoice_number"],
        issue_date=values["issue_date"],
        buyer_name=values["buyer_name"],
        buyer_tax_id=tax_ids[0] if tax_ids else None,
        seller_name=values["seller_name"],
        seller_tax_id=tax_ids[1] if len(tax_ids) > 1 else None,
        amount_excluding_tax=amount,
        tax_amount=tax,
        total_amount=values["total_amount"],
        confidence=round((len(values) - len(warnings)) / len(values), 2),
        warnings=warnings,
    )


def create_invoice_recognition_router() -> APIRouter:
    router = APIRouter()

    @router.post(
        "/v1/internal/invoices:recognize",
        response_model=InvoiceRecognitionResponse,
        include_in_schema=False,
    )
    async def recognize_invoices(
        files: Annotated[list[UploadFile], File(...)],
        internal_token: Annotated[str, Header(alias="X-Internal-Token")],
        tenant_id: Annotated[str, Header(alias="X-Tenant-Id")],
        user_id: Annotated[str, Header(alias="X-User-Id")],
    ) -> InvoiceRecognitionResponse:
        _check_token(internal_token)
        if not tenant_id.strip() or not user_id.strip():
            raise HTTPException(status_code=400, detail="Tenant and user are required.")
        if not files or len(files) > MAX_FILES:
            raise HTTPException(status_code=400, detail="Upload between 1 and 20 invoices.")
        invoices: list[InvoiceItem] = []
        for upload in files:
            file_name = upload.filename or "invoice.pdf"
            if not file_name.lower().endswith(".pdf"):
                raise HTTPException(status_code=415, detail="Only PDF is supported.")
            content = await upload.read(MAX_FILE_BYTES + 1)
            if not content or len(content) > MAX_FILE_BYTES:
                raise HTTPException(status_code=413, detail="Invoice is empty or too large.")
            text = _pdf_text(content)
            method = "PDF_TEXT"
            if len(re.sub(r"\s+", "", text)) < MIN_NATIVE_TEXT_LENGTH:
                text = await _ocr_text(file_name, content)
                method = "PADDLE_OCR"
            invoices.append(parse_invoice(
                file_name,
                text,
                method,
                hashlib.sha256(content).hexdigest(),
            ))

        def summed(field: str) -> Decimal:
            return sum(
                (getattr(item, field) or Decimal("0") for item in invoices),
                start=Decimal("0"),
            )

        return InvoiceRecognitionResponse(
            invoices=invoices,
            total_amount=summed("total_amount"),
            total_amount_excluding_tax=summed("amount_excluding_tax"),
            total_tax_amount=summed("tax_amount"),
            file_count=len(invoices),
            recognized_count=sum(not item.warnings for item in invoices),
        )

    return router
