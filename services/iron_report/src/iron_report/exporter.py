from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from scheduling.scheduler import SchedulingRuntimeOptions
from task_manager.artifact_service import resolve_artifact_path
from task_manager.pipeline.store import create_file_artifact, list_stage_runs
from task_manager.service import TaskManagerService

from iron_report.config import Settings
from iron_report.errors import IronReportError
from iron_report.schemas import IronReportTaskInput, IronReportTaskOutput, OutputFormat, ReportType


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


class ExportStateStore:
    def __init__(self, root: Path) -> None:
        self.root = Path(root).expanduser().resolve()

    def get(self, task_id: str) -> dict[str, Any]:
        path = self._path(task_id)
        if not path.is_file():
            return {}
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}

    def save(self, task_id: str, payload: dict[str, Any]) -> None:
        self.root.mkdir(parents=True, exist_ok=True)
        path = self._path(task_id)
        temporary = path.with_suffix(".tmp")
        content = {
            **payload,
            "task_id": task_id,
            "updated_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
        }
        temporary.write_text(json.dumps(content, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def _path(self, task_id: str) -> Path:
        safe = re.sub(r"[^A-Za-z0-9_-]", "_", task_id)
        return self.root / f"{safe}.json"


class IronReportExporter:
    def __init__(self, settings: Settings, options: SchedulingRuntimeOptions, state: ExportStateStore) -> None:
        self.settings = settings
        self.options = options
        self.state = state
        self.tasks = TaskManagerService(options)

    async def accept_task_result(self, payload: dict[str, Any]) -> dict[str, Any]:
        task_id = str(payload.get("task_id") or "").strip()
        run_id = str(payload.get("run_id") or "").strip()
        if not task_id or not run_id:
            raise IronReportError(
                "IRON_REPORT_RESULT_INVALID",
                "Framework 结果缺少 task_id 或 run_id",
                status_code=422,
            )
        task = await self.tasks.get_task(task_id)
        if task is None or task.task_type != "iron.report.generate" or task.current_run_id != run_id:
            raise IronReportError(
                "IRON_REPORT_RESULT_TASK_MISMATCH",
                "Framework 结果与铁矿石报告任务不匹配",
                status_code=409,
            )
        if str(payload.get("status") or "completed") != "completed":
            self.state.save(
                task_id,
                {
                    "status": "FAILED",
                    "stage": "FAILED",
                    "error_code": "IRON_REPORT_AGENT_FAILED",
                    "error_message": str(payload.get("error_message") or "报告 Agent 执行失败"),
                    "retryable": True,
                },
            )
            return {"status": "FAILED", "task_id": task_id}

        try:
            output = IronReportTaskOutput.model_validate(payload.get("output") or {})
            task_input = IronReportTaskInput.model_validate(task.input_payload_json)
            self._validate_result_dates(task_input, output)
            existing = await self.tasks.list_task_artifacts(task_id)
            chart_rows = [
                row for row in existing
                if str((row.metadata_json or {}).get("mime") or "").startswith("image/") and row.content_uri
            ]
            if len(chart_rows) < 2:
                raise IronReportError(
                    "IRON_REPORT_CHART_ARTIFACT_MISSING",
                    "报告 Agent 未生成两个有效图表产物",
                    status_code=422,
                    retryable=True,
                    details={"found": len(chart_rows), "required": 2},
                )
            chart_paths = [self._verified_artifact_path(row) for row in chart_rows[:2]]
            self.state.save(task_id, {"status": "RUNNING", "stage": "EXPORTING", "progress": 90})
            artifact_ids = await self._export(
                task_id=task_id,
                run_id=run_id,
                task_input=task_input,
                report=output,
                chart_paths=chart_paths,
                existing=existing,
            )
            self.state.save(
                task_id,
                {
                    "status": "COMPLETED",
                    "stage": "COMPLETED",
                    "progress": 100,
                    "artifact_ids": artifact_ids,
                },
            )
            return {"status": "COMPLETED", "task_id": task_id, "artifact_ids": artifact_ids}
        except IronReportError as exc:
            self.state.save(
                task_id,
                {
                    "status": "FAILED",
                    "stage": "FAILED",
                    "error_code": exc.code,
                    "error_message": str(exc),
                    "retryable": exc.retryable,
                    "details": exc.details,
                },
            )
            raise
        except Exception as exc:
            self.state.save(
                task_id,
                {
                    "status": "FAILED",
                    "stage": "FAILED",
                    "error_code": "IRON_REPORT_EXPORT_FAILED",
                    "error_message": str(exc),
                    "retryable": True,
                },
            )
            raise IronReportError(
                "IRON_REPORT_EXPORT_FAILED",
                "铁矿石报告导出失败",
                status_code=500,
                retryable=True,
            ) from exc

    async def _export(
        self,
        *,
        task_id: str,
        run_id: str,
        task_input: IronReportTaskInput,
        report: IronReportTaskOutput,
        chart_paths: list[Path],
        existing: list[Any],
    ) -> list[str]:
        stage_runs = await list_stage_runs(run_id)
        if not stage_runs:
            raise IronReportError(
                "IRON_REPORT_STAGE_NOT_FOUND",
                "报告任务缺少 Agent 阶段记录",
                status_code=500,
                retryable=True,
            )
        stage_run_id = stage_runs[-1].id
        work_dir = self.settings.export_state_root / "work" / task_id / run_id
        work_dir.mkdir(parents=True, exist_ok=True)
        base_name = _safe_base_name(report.title)
        docx_path = work_dir / f"{base_name}.docx"

        by_kind = {
            str((row.metadata_json or {}).get("kind") or ""): row
            for row in existing
            if row.content_uri
        }
        needs_docx = (
            OutputFormat.DOCX in task_input.output_formats and "DOCX" not in by_kind
        ) or (
            OutputFormat.PDF in task_input.output_formats and "PDF" not in by_kind
        )
        if needs_docx:
            self._write_docx(docx_path, report, chart_paths)
        artifact_ids: list[str] = []
        if OutputFormat.DOCX in task_input.output_formats:
            row = by_kind.get("DOCX") or await self._publish(
                docx_path,
                task_id=task_id,
                run_id=run_id,
                stage_run_id=stage_run_id,
                name=docx_path.name,
                mime=DOCX_MIME,
                kind="DOCX",
            )
            artifact_ids.append(row.id)

        if OutputFormat.PDF in task_input.output_formats:
            row = by_kind.get("PDF")
            if row is None:
                pdf_path = self._convert_pdf(docx_path, work_dir)
                row = await self._publish(
                    pdf_path,
                    task_id=task_id,
                    run_id=run_id,
                    stage_run_id=stage_run_id,
                    name=pdf_path.name,
                    mime=PDF_MIME,
                    kind="PDF",
                )
            artifact_ids.append(row.id)
        return artifact_ids

    def _write_docx(self, target: Path, report: IronReportTaskOutput, charts: list[Path]) -> None:
        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10.5)
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

        title = document.add_heading(report.title, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph(
            f"报告日期：{report.report_date.isoformat()}    数据截至：{report.data_as_of_date.isoformat()}"
        )
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_heading("核心摘要", level=1)
        document.add_paragraph(report.executive_summary)

        document.add_heading("核心指标", level=1)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["指标", "数值", "单位", "数据日期", "来源"]):
            cell.text = text
        for metric in report.metrics:
            cells = table.add_row().cells
            values = [metric.label, str(metric.value), metric.unit, metric.data_as_of_date.isoformat(), metric.source_id]
            for cell, value in zip(cells, values):
                cell.text = value

        for section in report.sections:
            document.add_heading(section.heading, level=1)
            document.add_paragraph(section.summary)
            for label, items in (("主要发现", section.findings), ("证据", section.evidence), ("风险", section.risks)):
                if not items:
                    continue
                document.add_paragraph(label, style="Heading 2")
                for item in items:
                    document.add_paragraph(item, style="List Bullet")

        document.add_heading("动态图表", level=1)
        for index, chart in enumerate(charts, start=1):
            document.add_picture(str(chart), width=Inches(6.2))
            caption = document.add_paragraph(f"图 {index}：由报告任务基于准备数据动态生成")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_heading("来源", level=1)
        for source in report.sources:
            line = f"[{source.source_id}] {source.title}"
            if source.published_at:
                line += f"（{source.published_at}）"
            if source.url:
                line += f" {source.url}"
            document.add_paragraph(line, style="List Bullet")

        document.add_heading("数据限制", level=1)
        for item in report.limitations:
            document.add_paragraph(item, style="List Bullet")
        document.add_paragraph(f"研究方式：{report.research_status}")
        document.add_paragraph(f"生成时间：{report.generated_at.isoformat()}")
        document.save(target)

    def _convert_pdf(self, docx_path: Path, output_dir: Path) -> Path:
        profile = output_dir / "libreoffice-profile"
        profile.mkdir(parents=True, exist_ok=True)
        command = [
            self.settings.libreoffice_path,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--nofirststartwizard",
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--convert-to",
            "pdf:writer_pdf_Export",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ]
        try:
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=self.settings.libreoffice_timeout_seconds,
                env={**os.environ, "HOME": str(output_dir)},
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise IronReportError(
                "IRON_REPORT_PDF_CONVERSION_FAILED",
                "LibreOffice PDF 转换未完成",
                status_code=500,
                retryable=True,
                details={"reason": str(exc)},
            ) from exc
        pdf_path = output_dir / f"{docx_path.stem}.pdf"
        if completed.returncode != 0 or not pdf_path.is_file() or pdf_path.stat().st_size == 0:
            raise IronReportError(
                "IRON_REPORT_PDF_CONVERSION_FAILED",
                "LibreOffice PDF 转换失败",
                status_code=500,
                retryable=True,
                details={
                    "exitCode": completed.returncode,
                    "stderr": completed.stderr[-2000:],
                    "stdout": completed.stdout[-2000:],
                },
            )
        return pdf_path

    async def _publish(
        self,
        source: Path,
        *,
        task_id: str,
        run_id: str,
        stage_run_id: str,
        name: str,
        mime: str,
        kind: str,
    ):
        if not source.is_file() or source.stat().st_size == 0:
            raise IronReportError(
                "IRON_REPORT_ARTIFACT_INVALID",
                f"{kind} 导出文件无效",
                status_code=500,
                retryable=True,
            )
        artifact_id = uuid.uuid4().hex
        date_path = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d")
        directory = self.options.local_python_artifact_dir / date_path / artifact_id
        directory.mkdir(parents=True, exist_ok=False)
        target = directory / name
        try:
            shutil.copyfile(source, target)
            checksum = _sha256(target)
            content_uri = target.resolve().relative_to(self.options.local_python_artifact_dir.resolve()).as_posix()
            return await create_file_artifact(
                artifact_id=artifact_id,
                task_id=task_id,
                run_id=run_id,
                stage_run_id=stage_run_id,
                content_uri=content_uri,
                checksum=checksum,
                summary=name,
                metadata={"name": name, "mime": mime, "kind": kind, "size": target.stat().st_size},
            )
        except Exception:
            shutil.rmtree(directory, ignore_errors=True)
            raise

    def _verified_artifact_path(self, row) -> Path:
        try:
            path = resolve_artifact_path(self.options.local_python_artifact_dir, row.content_uri)
        except ValueError as exc:
            raise IronReportError(
                "IRON_REPORT_ARTIFACT_PATH_INVALID", "图表产物路径无效", status_code=500
            ) from exc
        if not path.is_file() or _sha256(path) != row.checksum:
            raise IronReportError(
                "IRON_REPORT_ARTIFACT_CORRUPT", "图表产物校验失败", status_code=409
            )
        return path

    def _validate_result_dates(self, task_input: IronReportTaskInput, output: IronReportTaskOutput) -> None:
        if output.report_type != task_input.report_type:
            raise IronReportError("IRON_REPORT_RESULT_MISMATCH", "报告类型与任务不一致", status_code=422)
        if output.report_date != task_input.report_date or output.data_as_of_date != task_input.data_as_of_date:
            raise IronReportError(
                "IRON_REPORT_RESULT_DATE_MISMATCH",
                "报告日期或数据截至日期与任务不一致",
                status_code=422,
                details={
                    "expectedReportDate": task_input.report_date.isoformat(),
                    "actualReportDate": output.report_date.isoformat(),
                    "expectedDataAsOfDate": task_input.data_as_of_date.isoformat(),
                    "actualDataAsOfDate": output.data_as_of_date.isoformat(),
                },
            )


def _safe_base_name(title: str) -> str:
    normalized = re.sub(r"[\\/:*?\"<>|\r\n\t]", "_", title).strip(" ._")
    return normalized[:120] or "iron-ore-report"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
