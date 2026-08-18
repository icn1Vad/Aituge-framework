from __future__ import annotations

import hashlib
import html
import io
import json
import logging
import os
import re
import tempfile
import threading
import zipfile
from dataclasses import dataclass
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path
from typing import Any, Iterable, Mapping, Protocol

from fastapi import UploadFile
from pypdf import PdfReader, PdfWriter


SERVICE_NAME = "contract-ocr"
SERVICE_VERSION = "0.3.0"
_PDF_CONTENT_TYPE = "application/pdf"
_WORD_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_HTML_TAG = re.compile(r"<[^>]+>")
_ROW_TAG = re.compile(r"<tr[^>]*>(.*?)</tr>", re.IGNORECASE | re.DOTALL)
_CELL_TAG = re.compile(r"<t[dh][^>]*>(.*?)</t[dh]>", re.IGNORECASE | re.DOTALL)
logger = logging.getLogger(__name__)


class ContractOcrError(RuntimeError):
    def __init__(self, code: str, message: str, *, status_code: int, retryable: bool = False) -> None:
        super().__init__(message)
        self.code = code
        self.status_code = status_code
        self.retryable = retryable


@dataclass(frozen=True, slots=True)
class OcrSettings:
    max_file_bytes: int = 25 * 1024 * 1024
    max_pages: int = 200
    min_text_chars_per_page: int = 24
    max_output_bytes: int = 50 * 1024 * 1024
    structure_dpi: int = 300
    structure_max_width: int = 3800
    structure_tile_height: int = 3200
    structure_tile_overlap: int = 320

    @classmethod
    def from_environment(cls) -> "OcrSettings":
        defaults = cls()
        return cls(
            max_file_bytes=_positive_int("CONTRACT_OCR_MAX_FILE_BYTES", defaults.max_file_bytes),
            max_pages=_positive_int("CONTRACT_OCR_MAX_PAGES", defaults.max_pages),
            min_text_chars_per_page=_positive_int(
                "CONTRACT_OCR_MIN_TEXT_CHARS_PER_PAGE", defaults.min_text_chars_per_page
            ),
            max_output_bytes=_positive_int("CONTRACT_OCR_MAX_OUTPUT_BYTES", defaults.max_output_bytes),
            structure_dpi=_positive_int("CONTRACT_OCR_STRUCTURE_DPI", defaults.structure_dpi),
            structure_max_width=_positive_int(
                "CONTRACT_OCR_STRUCTURE_MAX_WIDTH", defaults.structure_max_width
            ),
            structure_tile_height=_positive_int(
                "CONTRACT_OCR_STRUCTURE_TILE_HEIGHT", defaults.structure_tile_height
            ),
            structure_tile_overlap=_positive_int(
                "CONTRACT_OCR_STRUCTURE_TILE_OVERLAP", defaults.structure_tile_overlap
            ),
        )


@dataclass(frozen=True, slots=True)
class PdfInspection:
    classification: str
    page_count: int
    text_page_count: int
    pages_requiring_ocr: tuple[int, ...]
    source_sha256: str


@dataclass(frozen=True, slots=True)
class OcrConversion:
    inspection: PdfInspection
    archive: bytes


@dataclass(frozen=True, slots=True)
class OcrStructure:
    inspection: PdfInspection
    pages: list[dict[str, Any]]


@dataclass(frozen=True, slots=True)
class RenderedTile:
    page_number: int
    page_count: int
    tile_index: int
    y_offset: int
    full_width: int
    full_height: int
    path: Path


class PdfPageRenderer(Protocol):
    def render_tiles(self, source_path: Path, workdir: Path) -> list[RenderedTile]:
        ...


class StructureEngine(Protocol):
    @property
    def initialized(self) -> bool:
        ...

    def parse_page(self, page_path: Path) -> Iterable[Mapping[str, Any]]:
        ...


class PaddleStructureEngine:
    """Lazily creates PP-StructureV3 so health checks never download model weights."""

    def __init__(self) -> None:
        self._pipeline: Any | None = None
        self._lock = threading.Lock()

    @property
    def initialized(self) -> bool:
        return self._pipeline is not None

    def parse_page(self, page_path: Path) -> Iterable[Mapping[str, Any]]:
        pipeline = self._get_pipeline()
        results = pipeline.predict(
            str(page_path),
            use_table_recognition=True,
            use_wired_table_cells_trans_to_html=True,
            use_wireless_table_cells_trans_to_html=True,
            format_block_content=False,
        )
        return [_paddle_result_payload(result) for result in results]

    def _get_pipeline(self) -> Any:
        if self._pipeline is not None:
            return self._pipeline
        with self._lock:
            if self._pipeline is None:
                try:
                    from paddleocr import PPStructureV3

                    self._pipeline = PPStructureV3(lang="ch", use_table_recognition=True)
                except Exception as exc:  # Paddle model download and initialization are external work.
                    raise ContractOcrError(
                        "OCR_MODEL_INITIALIZATION_FAILED",
                        "OCR模型初始化失败",
                        status_code=503,
                        retryable=True,
                    ) from exc
        return self._pipeline


class PdfiumPageRenderer:
    """Renders high-resolution page images and vertically tiles oversized pages."""

    def __init__(self, settings: OcrSettings) -> None:
        self.settings = settings

    def render_tiles(self, source_path: Path, workdir: Path) -> list[RenderedTile]:
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(source_path))
        rendered: list[RenderedTile] = []
        try:
            page_count = len(document)
            for page_index in range(page_count):
                page = document[page_index]
                try:
                    page_width, _ = page.get_size()
                    requested_scale = self.settings.structure_dpi / 72.0
                    width_scale = self.settings.structure_max_width / max(float(page_width), 1.0)
                    scale = min(requested_scale, width_scale)
                    bitmap = page.render(scale=scale)
                    try:
                        image = bitmap.to_pil().convert("RGB")
                    finally:
                        bitmap.close()
                    try:
                        rendered.extend(
                            self._save_tiles(
                                image,
                                workdir,
                                page_number=page_index + 1,
                                page_count=page_count,
                            )
                        )
                    finally:
                        image.close()
                finally:
                    page.close()
        finally:
            document.close()
        return rendered

    def _save_tiles(
        self,
        image: Any,
        workdir: Path,
        *,
        page_number: int,
        page_count: int,
    ) -> list[RenderedTile]:
        width, height = image.size
        tile_height = min(self.settings.structure_tile_height, height)
        overlap = min(self.settings.structure_tile_overlap, max(0, tile_height // 3))
        step = max(1, tile_height - overlap)
        offsets = [0] if height <= tile_height else list(range(0, height, step))
        if offsets[-1] + tile_height < height:
            offsets.append(height - tile_height)
        offsets = sorted(set(min(offset, max(0, height - tile_height)) for offset in offsets))
        tiles: list[RenderedTile] = []
        for tile_index, y_offset in enumerate(offsets):
            bottom = min(height, y_offset + tile_height)
            tile_image = image.crop((0, y_offset, width, bottom))
            tile_path = workdir / f"page-{page_number}-tile-{tile_index}.png"
            try:
                tile_image.save(tile_path, format="PNG")
            finally:
                tile_image.close()
            tiles.append(
                RenderedTile(
                    page_number=page_number,
                    page_count=page_count,
                    tile_index=tile_index,
                    y_offset=y_offset,
                    full_width=width,
                    full_height=height,
                    path=tile_path,
                )
            )
        return tiles


class ContractOcrService:
    service_name = SERVICE_NAME
    service_version = SERVICE_VERSION

    def __init__(
        self,
        settings: OcrSettings,
        engine: StructureEngine | None = None,
        renderer: PdfPageRenderer | None = None,
    ) -> None:
        self.settings = settings
        self.engine = engine or PaddleStructureEngine()
        self.renderer = renderer or PdfiumPageRenderer(settings)

    @classmethod
    def from_environment(cls) -> "ContractOcrService":
        return cls(OcrSettings.from_environment())

    @property
    def model_initialized(self) -> bool:
        return self.engine.initialized

    async def read_pdf_upload(self, upload: UploadFile) -> bytes:
        content_type = (upload.content_type or "").lower()
        name = (upload.filename or "").lower()
        if content_type not in {_PDF_CONTENT_TYPE, "application/octet-stream", ""} and not name.endswith(".pdf"):
            raise ContractOcrError("FILE_TYPE_UNSUPPORTED", "OCR服务仅支持PDF合同", status_code=415)
        content = await upload.read(self.settings.max_file_bytes + 1)
        if not content:
            raise ContractOcrError("FILE_EMPTY", "PDF合同不能为空", status_code=422)
        if len(content) > self.settings.max_file_bytes:
            raise ContractOcrError("FILE_TOO_LARGE", "PDF合同超过OCR处理大小限制", status_code=413)
        if not content.startswith(b"%PDF-"):
            raise ContractOcrError("PDF_INVALID", "上传内容不是有效PDF", status_code=422)
        return content

    def runtime_versions(self, *, verify_import: bool) -> tuple[str, str]:
        paddleocr_version = _package_version("paddleocr")
        paddlepaddle_version = _package_version("paddlepaddle")
        if verify_import:
            try:
                import paddle  # noqa: F401
                import paddleocr  # noqa: F401
            except Exception as exc:
                raise ContractOcrError(
                    "OCR_RUNTIME_UNAVAILABLE", "PaddleOCR运行时不可用", status_code=503, retryable=True
                ) from exc
        return paddleocr_version, paddlepaddle_version

    def inspect_pdf(self, content: bytes) -> PdfInspection:
        reader = _open_pdf(content)
        page_count = len(reader.pages)
        if page_count == 0:
            raise ContractOcrError("PDF_INVALID", "PDF不包含有效页面", status_code=422)
        if page_count > self.settings.max_pages:
            raise ContractOcrError("PDF_PAGE_LIMIT_EXCEEDED", "PDF页数超过OCR处理限制", status_code=413)
        pages_requiring_ocr: list[int] = []
        text_page_count = 0
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                raise ContractOcrError(
                    "PDF_TEXT_EXTRACTION_FAILED", f"PDF第{page_number}页文本层读取失败", status_code=422
                ) from exc
            if _has_usable_text(text, self.settings.min_text_chars_per_page):
                text_page_count += 1
            else:
                pages_requiring_ocr.append(page_number)
        classification = (
            "NATIVE_TEXT"
            if not pages_requiring_ocr
            else "OCR_REQUIRED"
            if text_page_count == 0
            else "MIXED"
        )
        return PdfInspection(
            classification=classification,
            page_count=page_count,
            text_page_count=text_page_count,
            pages_requiring_ocr=tuple(pages_requiring_ocr),
            source_sha256=_sha256(content),
        )

    def convert_pdf(self, file_name: str, content: bytes) -> OcrConversion:
        inspection = self.inspect_pdf(content)
        if inspection.classification == "NATIVE_TEXT":
            pages = _extract_text_layer_pages(content, inspection)
            engine = {"name": "PDF_TEXT_LAYER", "pypdf": _package_version("pypdf")}
        else:
            with tempfile.TemporaryDirectory(prefix="contract-ocr-") as directory:
                workdir = Path(directory)
                source_path = workdir / "source.pdf"
                source_path.write_bytes(content)
                pages = self._recognize_pages(source_path, inspection)
            engine = {"name": "PP-StructureV3", "paddleocr": _package_version("paddleocr")}

        docx = _build_docx(pages)
        manifest = {
            "schema_version": "1.0",
            "source_file_name": file_name,
            "source_sha256": inspection.source_sha256,
            "classification": inspection.classification,
            "page_count": inspection.page_count,
            "pages_requiring_ocr": list(inspection.pages_requiring_ocr),
            "engine": engine,
            "pages": pages,
        }
        archive = _build_archive(docx, manifest)
        if len(archive) > self.settings.max_output_bytes:
            raise ContractOcrError("OCR_OUTPUT_TOO_LARGE", "OCR转换结果超过大小限制", status_code=413)
        return OcrConversion(inspection=inspection, archive=archive)

    def structure_pdf(self, content: bytes) -> OcrStructure:
        inspection = self.inspect_pdf(content)
        with tempfile.TemporaryDirectory(prefix="contract-ocr-structure-") as directory:
            workdir = Path(directory)
            source_path = workdir / "source.pdf"
            source_path.write_bytes(content)
            tiles = self.renderer.render_tiles(source_path, workdir)
            pages_by_number: dict[int, dict[str, Any]] = {}
            for tile in tiles:
                try:
                    results = list(self.engine.parse_page(tile.path))
                except ContractOcrError:
                    raise
                except Exception as exc:
                    logger.exception(
                        "Structured OCR inference failed for page %s tile %s",
                        tile.page_number,
                        tile.tile_index,
                    )
                    raise ContractOcrError(
                        "OCR_INFERENCE_FAILED",
                        f"PDF第{tile.page_number}页结构识别失败",
                        status_code=503,
                        retryable=True,
                    ) from exc
                if not results:
                    raise ContractOcrError(
                        "OCR_NO_RESULT",
                        f"PDF第{tile.page_number}页没有结构识别结果",
                        status_code=422,
                    )
                page = pages_by_number.setdefault(
                    tile.page_number,
                    {
                        "page_number": tile.page_number,
                        "page_count": tile.page_count,
                        "width": tile.full_width,
                        "height": tile.full_height,
                        "blocks": [],
                        "tables": [],
                    },
                )
                for payload in results:
                    segment = _manifest_page(payload, tile.page_number, tile.page_count)
                    for block in segment["blocks"]:
                        translated = dict(block)
                        translated["bbox"] = _offset_boxes(translated.get("bbox"), tile.y_offset)
                        translated["tile_index"] = tile.tile_index
                        page["blocks"].append(translated)
                    for table in segment["tables"]:
                        translated_table = _offset_table(table, tile.y_offset)
                        translated_table["tile_index"] = tile.tile_index
                        page["tables"].append(translated_table)
            pages = [pages_by_number[number] for number in sorted(pages_by_number)]
        return OcrStructure(inspection=inspection, pages=pages)

    def _recognize_pages(self, source_path: Path, inspection: PdfInspection) -> list[dict[str, Any]]:
        reader = _open_pdf(source_path.read_bytes())
        pages: list[dict[str, Any]] = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_path = source_path.parent / f"page-{page_number}.pdf"
            writer = PdfWriter()
            writer.add_page(page)
            with page_path.open("wb") as output:
                writer.write(output)
            try:
                results = list(self.engine.parse_page(page_path))
            except ContractOcrError:
                raise
            except Exception as exc:
                # Keep the user-facing error stable, but retain the root cause
                # server-side without logging document bytes or extracted text.
                logger.exception("OCR inference failed for page %s", page_number)
                raise ContractOcrError(
                    "OCR_INFERENCE_FAILED", f"PDF第{page_number}页OCR失败", status_code=503, retryable=True
                ) from exc
            if not results:
                raise ContractOcrError(
                    "OCR_NO_RESULT", f"PDF第{page_number}页没有OCR结果", status_code=422
                )
            # A single-page input normally yields one result. Keep every result to avoid silently discarding output.
            for payload in results:
                pages.append(_manifest_page(payload, page_number, inspection.page_count))
        return pages


def _positive_int(name: str, default: int) -> int:
    raw = os.getenv(name)
    if raw is None or not raw.strip():
        return default
    try:
        value = int(raw)
    except ValueError:
        return default
    return value if value > 0 else default


def _package_version(name: str) -> str:
    try:
        return version(name)
    except PackageNotFoundError as exc:
        raise ContractOcrError("OCR_RUNTIME_UNAVAILABLE", f"缺少运行依赖：{name}", status_code=503) from exc


def _open_pdf(content: bytes) -> PdfReader:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ContractOcrError("PDF_INVALID", "PDF无法解析", status_code=422) from exc
    if reader.is_encrypted:
        raise ContractOcrError("PDF_ENCRYPTED", "暂不支持加密PDF合同", status_code=422)
    return reader


def _has_usable_text(text: str, minimum: int) -> bool:
    normalized = re.sub(r"\s+", "", text)
    if len(normalized) < minimum:
        return False
    visible = sum(character.isalnum() or "\u4e00" <= character <= "\u9fff" for character in normalized)
    return visible / max(1, len(normalized)) >= 0.45


def _sha256(content: bytes) -> str:
    return "sha256:" + hashlib.sha256(content).hexdigest()


def _paddle_result_payload(result: Any) -> Mapping[str, Any]:
    payload = getattr(result, "json", result)
    if callable(payload):
        payload = payload()
    if isinstance(payload, Mapping) and isinstance(payload.get("res"), Mapping):
        payload = payload["res"]
    if not isinstance(payload, Mapping):
        raise ContractOcrError("OCR_RESULT_INVALID", "OCR服务返回了无效结果", status_code=502, retryable=True)
    return _json_value(payload)


def _json_value(value: Any) -> Any:
    if isinstance(value, Mapping):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_json_value(item) for item in value]
    if hasattr(value, "tolist"):
        return _json_value(value.tolist())
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _extract_text_layer_pages(
    content: bytes,
    inspection: PdfInspection,
) -> list[dict[str, Any]]:
    reader = _open_pdf(content)
    pages: list[dict[str, Any]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        try:
            extracted = page.extract_text() or ""
        except Exception as exc:
            raise ContractOcrError(
                "PDF_TEXT_EXTRACTION_FAILED",
                f"PDF第{page_number}页文本层读取失败",
                status_code=422,
            ) from exc
        lines = [_clean_text(line) for line in extracted.splitlines()]
        blocks = [
            {
                "id": f"text-{page_number}-{index}",
                "order": index,
                "label": "text",
                "text": line,
                "bbox": [],
            }
            for index, line in enumerate((line for line in lines if line), start=1)
        ]
        if not blocks:
            raise ContractOcrError(
                "PDF_TEXT_EXTRACTION_FAILED",
                f"PDF第{page_number}页没有可转换的文字层",
                status_code=422,
            )
        pages.append(
            {
                "page_number": page_number,
                "page_count": inspection.page_count,
                "width": float(page.mediabox.width),
                "height": float(page.mediabox.height),
                "blocks": blocks,
                "tables": [],
            }
        )
    return pages


def _manifest_page(payload: Mapping[str, Any], page_number: int, page_count: int) -> dict[str, Any]:
    raw_blocks = payload.get("parsing_res_list")
    blocks: list[dict[str, Any]] = []
    if isinstance(raw_blocks, list):
        for index, raw in enumerate(raw_blocks, start=1):
            item = raw if isinstance(raw, Mapping) else {}
            blocks.append(
                {
                    "id": item.get("block_id") or item.get("index") or index,
                    "order": item.get("block_order") or item.get("order_index") or index,
                    "label": item.get("block_label") or item.get("label") or "text",
                    "text": item.get("block_content") or item.get("content") or "",
                    "bbox": item.get("block_bbox") or item.get("bbox") or [],
                }
            )
    tables = payload.get("table_res_list")
    return {
        "page_number": page_number,
        "page_count": page_count,
        "width": payload.get("width"),
        "height": payload.get("height"),
        "blocks": _json_value(blocks),
        "tables": _json_value(tables if isinstance(tables, list) else []),
    }


def _build_docx(pages: list[Mapping[str, Any]]) -> bytes:
    from docx import Document

    document = Document()
    body_count = 0
    for page in pages:
        raw_blocks = page.get("blocks")
        blocks = [item for item in raw_blocks if isinstance(item, Mapping)] if isinstance(raw_blocks, list) else []
        blocks.sort(key=lambda block: int(block.get("order") or 0))
        for block in blocks:
            text = str(block.get("text") or "").strip()
            label = str(block.get("label") or "text").lower()
            if not text:
                continue
            if label == "table" and _append_html_table(document, text):
                body_count += 1
                continue
            cleaned = _clean_text(text)
            if not cleaned:
                continue
            if label in {"title", "doc_title", "section_title"}:
                document.add_heading(cleaned, level=1)
            elif label in {"paragraph_title", "sub_title", "heading"}:
                document.add_heading(cleaned, level=2)
            else:
                for paragraph in cleaned.splitlines():
                    normalized = paragraph.strip()
                    if normalized:
                        document.add_paragraph(normalized)
            body_count += 1
    if body_count == 0:
        raise ContractOcrError("OCR_NO_TEXT", "OCR未识别到可审查文本", status_code=422)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _append_html_table(document: Any, value: str) -> bool:
    rows = []
    for raw_row in _ROW_TAG.findall(value):
        cells = [_clean_text(cell) for cell in _CELL_TAG.findall(raw_row)]
        if cells:
            rows.append(cells)
    if not rows:
        return False
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row):
            table.cell(row_index, column_index).text = cell
    return True


def _clean_text(value: str) -> str:
    return html.unescape(_HTML_TAG.sub("", value)).replace("\u00a0", " ").strip()


def _offset_boxes(value: Any, y_offset: int) -> Any:
    converted = _json_value(value)
    if (
        isinstance(converted, list)
        and len(converted) >= 4
        and all(isinstance(item, (int, float)) for item in converted[:4])
    ):
        result = list(converted)
        result[1] += y_offset
        result[3] += y_offset
        return result
    if isinstance(converted, list):
        return [_offset_boxes(item, y_offset) for item in converted]
    return converted


def _offset_table(value: Any, y_offset: int) -> dict[str, Any]:
    table = dict(_json_value(value)) if isinstance(value, Mapping) else {}
    for key in ("cell_box_list", "table_box", "bbox"):
        if key in table:
            table[key] = _offset_boxes(table[key], y_offset)
    prediction = table.get("table_ocr_pred")
    if isinstance(prediction, Mapping):
        translated_prediction = dict(prediction)
        if "rec_boxes" in translated_prediction:
            translated_prediction["rec_boxes"] = _offset_boxes(
                translated_prediction["rec_boxes"], y_offset
            )
        table["table_ocr_pred"] = translated_prediction
    return table


def _build_archive(docx: bytes, manifest: Mapping[str, Any]) -> bytes:
    output = io.BytesIO()
    try:
        with zipfile.ZipFile(output, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("contract.docx", docx)
            archive.writestr(
                "manifest.json",
                json.dumps(manifest, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
            )
    except Exception as exc:
        raise ContractOcrError("OCR_ARCHIVE_FAILED", "OCR结果打包失败", status_code=500) from exc
    return output.getvalue()
