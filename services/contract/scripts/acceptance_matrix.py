from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import time
import uuid
from collections import Counter
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document


TERMINAL_STATUSES = {"SUCCEEDED", "FAILED", "CANCELLED"}
RISK_CATEGORIES = {
    "RIGHTS_OBLIGATIONS_IMBALANCE",
    "PAYMENT",
    "BREACH",
    "LIABILITY",
    "TERMINATION",
    "MISSING_CLAUSE",
    "AMBIGUITY",
    "INTERNAL_CONFLICT",
}


def _call(
    url: str,
    *,
    method: str = "GET",
    headers: dict[str, str] | None = None,
    body: bytes | None = None,
) -> tuple[int, dict]:
    request = Request(url, data=body, headers=headers or {}, method=method)
    try:
        with urlopen(request, timeout=60) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _text_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 10 Tf 54 738 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode("ascii"))
        content.extend(value)
        content.extend(b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(content)


def _risk_docx() -> bytes:
    document = Document()
    document.add_heading("Service Agreement", level=1)
    document.add_paragraph("Party A (Client): Acme Company")
    document.add_paragraph("Party B (Provider): Beta Company")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Commercial term"
    table.rows[0].cells[1].text = "Agreement"
    for key, value in (
        ("Fees", "Beta Company shall pay the entire annual fee before services begin."),
        ("Acceptance", "Acme Company alone decides whether the services are accepted."),
        ("Refund", "All fees paid by Beta Company are non-refundable in every circumstance."),
    ):
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    document.add_heading("Liability and termination", level=2)
    document.add_paragraph(
        "Beta Company shall indemnify Acme Company for all losses without any monetary limit."
    )
    document.add_paragraph(
        "Acme Company may terminate this Agreement at any time without cause, while Beta Company "
        "may not terminate before the end of the five-year term."
    )
    document.add_heading("Conflicting payment dates", level=2)
    document.add_paragraph("Payment is due within ten days after invoice receipt.")
    document.add_paragraph("Payment is also stated to be due within sixty days after invoice receipt.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _multipart(
    *,
    filename: str,
    content_type: str,
    content: bytes,
    payload: dict,
) -> tuple[bytes, str]:
    boundary = f"contract-acceptance-{uuid.uuid4().hex}"
    parts = [
        ("file", filename, content_type, content),
        (
            "request",
            None,
            "application/json",
            json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
        ),
    ]
    body = bytearray()
    for name, part_filename, part_type, part_content in parts:
        body.extend(f"--{boundary}\r\n".encode())
        disposition = f'Content-Disposition: form-data; name="{name}"'
        if part_filename is not None:
            disposition += f'; filename="{part_filename}"'
        body.extend(f"{disposition}\r\n".encode())
        body.extend(f"Content-Type: {part_type}\r\n\r\n".encode())
        body.extend(part_content)
        body.extend(b"\r\n")
    body.extend(f"--{boundary}--\r\n".encode())
    return bytes(body), boundary


def _headers(token: str, marker: str, request_suffix: str) -> dict[str, str]:
    return {
        "X-Internal-Service": "continew-java",
        "X-Internal-Token": token,
        "X-User-Id": f"acceptance-user-{marker}",
        "X-Tenant-Id": "0",
        "X-Request-Id": f"acceptance-request-{marker}-{request_suffix}",
        "Idempotency-Key": f"acceptance-idempotency-{marker}-{request_suffix}",
    }


def _create_review(
    *,
    base_url: str,
    token: str,
    marker: str,
    suffix: str,
    contract_version_id: str,
    perspective: str,
    our_party_name: str,
    filename: str,
    content_type: str,
    content: bytes,
) -> tuple[dict, dict[str, str]]:
    headers = _headers(token, marker, suffix)
    payload = {
        "business_task_id": f"acceptance-task-{marker}-{suffix}",
        "contract_version_id": contract_version_id,
        "perspective": perspective,
        "our_party_name": our_party_name,
        "contract_type": "AUTO",
        "review_attitude": "NEUTRAL",
        "schema_version": "1.0",
    }
    body, boundary = _multipart(
        filename=filename,
        content_type=content_type,
        content=content,
        payload=payload,
    )
    create_headers = dict(headers)
    create_headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    code, response = _call(
        f"{base_url}/v1/contract-reviews",
        method="POST",
        headers=create_headers,
        body=body,
    )
    assert code == 201, response
    created = response["data"]
    print(
        "CREATED "
        f"scenario={suffix} review_id={created['review_id']} "
        f"document_id={created['document_id']} task_id={created.get('framework_task_id')}"
    )
    return created, headers


def _wait_for_success(
    *,
    base_url: str,
    review_id: str,
    headers: dict[str, str],
    timeout_seconds: int,
) -> dict:
    deadline = time.monotonic() + timeout_seconds
    last_snapshot: tuple[str, str | None, int | None] | None = None
    while True:
        code, response = _call(f"{base_url}/v1/contract-reviews/{review_id}", headers=headers)
        assert code == 200, response
        status = response["data"]
        snapshot = (
            status["status"],
            status.get("current_stage"),
            status.get("framework_attempt_no"),
        )
        if snapshot != last_snapshot:
            print(
                f"STATUS review_id={review_id} status={snapshot[0]} "
                f"stage={snapshot[1]} attempt={snapshot[2]}"
            )
            last_snapshot = snapshot
        if status["status"] in TERMINAL_STATUSES:
            assert status["status"] == "SUCCEEDED", status
            code, result = _call(
                f"{base_url}/v1/contract-reviews/{review_id}/result",
                headers=headers,
            )
            assert code == 200, result
            return result["data"]
        if time.monotonic() >= deadline:
            raise TimeoutError(f"Review {review_id} did not finish within {timeout_seconds}s")
        time.sleep(3)


def _validate_result(
    result: dict,
    *,
    perspective: str,
    our_party: str,
    counterparty: str,
) -> None:
    assert result["schema_version"] == "1.0"
    assert result["relationships"] == []
    assert result["result_hash"].startswith("sha256:") and len(result["result_hash"]) == 71
    profile = result["contract_profile"]
    assert profile["perspective"] == perspective
    assert profile["our_party"] == our_party
    assert profile["counterparty"] == counterparty
    assert profile["review_attitude"] == "NEUTRAL"

    findings = result["findings"]
    evidences = result["evidences"]
    findings_by_id = {item["finding_id"]: item for item in findings}
    evidences_by_id = {item["evidence_id"]: item for item in evidences}
    assert len(findings_by_id) == len(findings)
    assert len(evidences_by_id) == len(evidences)
    assert Counter(item["risk_level"] for item in findings) == Counter(
        {
            "HIGH": result["summary"]["high_count"],
            "MEDIUM": result["summary"]["medium_count"],
            "LOW": result["summary"]["low_count"],
            "INFO": result["summary"]["info_count"],
        }
    )
    for finding in findings:
        assert finding["perspective"] == perspective
        assert finding["our_party"] == our_party
        assert finding["counterparty"] == counterparty
        assert finding["evidence_ids"]
        assert all(evidence_id in evidences_by_id for evidence_id in finding["evidence_ids"])
    for evidence in evidences:
        assert evidence["finding_id"] in findings_by_id
        assert evidence["bounding_boxes"] == []
        if evidence["evidence_type"] == "ABSENCE":
            for field in (
                "block_id",
                "page_number",
                "char_start",
                "char_end",
                "quoted_text",
                "quoted_text_hash",
            ):
                assert evidence[field] is None
            assert evidence["checked_scope"] and evidence["verification_note"]
        else:
            quote = evidence["quoted_text"]
            assert quote
            assert evidence["char_end"] - evidence["char_start"] == len(quote)
            expected_hash = "sha256:" + hashlib.sha256(quote.encode("utf-8")).hexdigest()
            assert evidence["quoted_text_hash"] == expected_hash


def _dual_perspective(base_url: str, token: str, marker: str, timeout_seconds: int) -> None:
    content = _text_pdf(
        "Party A: Acme Company. Party B: Beta Company. Beta Company shall pay all fees in advance. "
        "Acme Company may terminate at any time without cause. Beta Company may not terminate early. "
        "Beta Company indemnifies Acme Company for all losses without a liability cap. "
        "Acme Company liability is capped at one hundred dollars."
    )
    version_id = f"acceptance-version-{marker}-dual"
    created_a, headers_a = _create_review(
        base_url=base_url,
        token=token,
        marker=marker,
        suffix="dual-party-a",
        contract_version_id=version_id,
        perspective="PARTY_A",
        our_party_name="Acme Company",
        filename="dual-perspective.pdf",
        content_type="application/pdf",
        content=content,
    )
    result_a = _wait_for_success(
        base_url=base_url,
        review_id=created_a["review_id"],
        headers=headers_a,
        timeout_seconds=timeout_seconds,
    )
    _validate_result(
        result_a,
        perspective="PARTY_A",
        our_party="Acme Company",
        counterparty="Beta Company",
    )

    created_b, headers_b = _create_review(
        base_url=base_url,
        token=token,
        marker=marker,
        suffix="dual-party-b",
        contract_version_id=version_id,
        perspective="PARTY_B",
        our_party_name="Beta Company",
        filename="dual-perspective.pdf",
        content_type="application/pdf",
        content=content,
    )
    assert created_b["document_id"] == created_a["document_id"]
    result_b = _wait_for_success(
        base_url=base_url,
        review_id=created_b["review_id"],
        headers=headers_b,
        timeout_seconds=timeout_seconds,
    )
    _validate_result(
        result_b,
        perspective="PARTY_B",
        our_party="Beta Company",
        counterparty="Acme Company",
    )
    assert result_a["contract_profile"] != result_b["contract_profile"]
    print("SCENARIO=DUAL_PERSPECTIVE/PASSED")


def _docx_risk(base_url: str, token: str, marker: str, timeout_seconds: int) -> None:
    created, headers = _create_review(
        base_url=base_url,
        token=token,
        marker=marker,
        suffix="docx-risk",
        contract_version_id=f"acceptance-version-{marker}-docx",
        perspective="PARTY_B",
        our_party_name="Beta Company",
        filename="table-risk.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        content=_risk_docx(),
    )
    result = _wait_for_success(
        base_url=base_url,
        review_id=created["review_id"],
        headers=headers,
        timeout_seconds=timeout_seconds,
    )
    _validate_result(
        result,
        perspective="PARTY_B",
        our_party="Beta Company",
        counterparty="Acme Company",
    )
    categories = {item["category"] for item in result["findings"]}
    assert categories & RISK_CATEGORIES, result
    print(
        f"SCENARIO=DOCX_RISK/PASSED findings={len(result['findings'])} "
        f"evidences={len(result['evidences'])}"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:19200")
    parser.add_argument("--timeout-seconds", type=int, default=900)
    parser.add_argument(
        "--scenario",
        choices=("all", "dual-perspective", "docx-risk"),
        default="all",
    )
    args = parser.parse_args()
    token = os.environ["CONTRACT_INTERNAL_TOKEN"]
    marker = uuid.uuid4().hex
    print(f"ACCEPTANCE_MARKER={marker}")
    if args.scenario in {"all", "dual-perspective"}:
        _dual_perspective(args.base_url, token, marker, args.timeout_seconds)
    if args.scenario in {"all", "docx-risk"}:
        _docx_risk(args.base_url, token, marker, args.timeout_seconds)
    print("ACCEPTANCE_MATRIX=PASSED")


if __name__ == "__main__":
    main()
