from __future__ import annotations

from io import BytesIO
import json

from docx import Document
from fastapi.testclient import TestClient

from contract.window_inspector import app


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("技术服务合同", level=1)
    document.add_paragraph("第一条 服务内容")
    document.add_paragraph("乙方应当按照约定完成系统部署。")
    document.add_paragraph("第二条 验收")
    document.add_paragraph("甲方应在收到交付物后五个工作日内组织验收。")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_window_inspector_renders_test_page() -> None:
    response = TestClient(app).get("/")

    assert response.status_code == 200
    assert "合同 IR Window 阶段验收" in response.text
    assert "并发 10 抽取全部 Window 并合并 IR" in response.text
    assert "模拟正式链路已有的 resolve_parties 结果" in response.text
    assert 'id="party-a-name"' in response.text
    assert 'id="party-b-name"' in response.text
    assert 'value="PARTY_A"' in response.text
    assert "Shadow Compare 只比较同类别的原文 Anchor" in response.text
    assert 'id="legacy-ir-file"' in response.text
    assert 'id="stage5-load"' in response.text


def test_window_inspector_reads_only_mounted_stage5_artifacts(tmp_path, monkeypatch) -> None:
    summary_path = tmp_path / "summary.json"
    result_path = tmp_path / "result.json"
    summary_path.write_text(json.dumps({"status": "SUCCEEDED", "finding_count": 1}), "utf-8")
    result_path.write_text(
        json.dumps({"success": True, "data": {"findings": [], "evidences": []}}),
        "utf-8",
    )
    monkeypatch.setenv("CONTRACT_STAGE5_SUMMARY_FILE", str(summary_path))
    monkeypatch.setenv("CONTRACT_STAGE5_RESULT_FILE", str(result_path))

    response = TestClient(app).get("/api/stage5-result")

    assert response.status_code == 200
    assert response.json() == {
        "summary": {"status": "SUCCEEDED", "finding_count": 1},
        "result": {"findings": [], "evidences": []},
    }


def test_window_inspector_returns_sections_windows_and_exact_coverage() -> None:
    response = TestClient(app).post(
        "/api/inspect",
        files={
            "file": (
                "技术服务合同.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["file_type"] == "docx"
    assert body["document_id"].startswith("window-document-")
    assert body["generation_id"].startswith("window-inspector-")
    assert body["block_count"] == 5
    assert body["section_count"] >= 3
    assert body["window_count"] >= 1
    assert body["coverage"] == {
        "valid": True,
        "expected_block_count": 5,
        "covered_block_count": 5,
        "split_block_ids": [],
        "missing_block_ids": [],
        "overlap_block_ids": [],
    }
    assert body["windows"][0]["offset_map"]
    assert len(body["expected_blocks"]) == body["block_count"]
    assert "第一条 服务内容" in "\n".join(item["source_text"] for item in body["windows"])


def test_window_inspector_rejects_unsupported_file_type() -> None:
    response = TestClient(app).post(
        "/api/inspect",
        files={"file": ("contract.txt", b"not supported", "text/plain")},
    )

    assert response.status_code == 415
    assert response.json()["detail"] == "测试页面仅接受 PDF 或 DOCX"
