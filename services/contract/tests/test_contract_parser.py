from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader, PdfWriter

from contract.errors import ContractError
from contract.parser.native import NativeContractParser
from pdf_factory import text_pdf_bytes


def test_pdf_parser_preserves_page_and_stable_block_identity(tmp_path: Path) -> None:
    path = tmp_path / "contract.pdf"
    path.write_bytes(text_pdf_bytes("Payment is due in 30 days."))
    parser = NativeContractParser()

    first = parser.parse(path, generation_id="generation-1")
    repeated = parser.parse(path, generation_id="generation-1")
    another_generation = parser.parse(path, generation_id="generation-2")

    assert first.page_count == 1
    assert first.blocks[0].page_number == 1
    assert first.blocks[0].text == "Payment is due in 30 days."
    assert first.blocks[0].block_id == repeated.blocks[0].block_id
    assert first.blocks[0].block_id != another_generation.blocks[0].block_id


def test_docx_parser_keeps_headings_manual_breaks_and_table_rows(tmp_path: Path) -> None:
    path = tmp_path / "contract.docx"
    document = Document()
    document.add_heading("第一章 服务内容", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("第一条 甲方委托乙方提供服务。")
    paragraph.add_run().add_break()
    paragraph.add_run("第二条 甲方应按期付款。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "付款节点"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "验收后"
    table.cell(1, 1).text = "100元"
    document.save(path)

    parsed = NativeContractParser().parse(path, generation_id="generation-docx")

    assert [item.block_type for item in parsed.blocks] == [
        "heading",
        "article",
        "article",
        "table_row",
        "table_row",
    ]
    assert parsed.blocks[1].heading_path == ["第一章 服务内容"]
    assert parsed.blocks[-1].metadata == {
        "table_no": 1,
        "row_no": 2,
        "leading_whitespace": 0,
    }
    assert all(item.char_end - item.char_start == len(item.text) for item in parsed.blocks)


def test_docx_parser_preserves_direct_and_style_native_numbering_metadata(tmp_path: Path) -> None:
    def attach_num_pr(properties, num_id: int, level: int) -> None:
        number_properties = OxmlElement("w:numPr")
        level_node = OxmlElement("w:ilvl")
        level_node.set(qn("w:val"), str(level))
        number_node = OxmlElement("w:numId")
        number_node.set(qn("w:val"), str(num_id))
        number_properties.append(level_node)
        number_properties.append(number_node)
        properties.append(number_properties)

    path = tmp_path / "numbered.docx"
    document = Document()
    direct = document.add_paragraph("原生直接编号条款")
    attach_num_pr(direct._p.get_or_add_pPr(), 17, 1)
    style = document.styles.add_style("ContractNativeList", WD_STYLE_TYPE.PARAGRAPH)
    attach_num_pr(style._element.get_or_add_pPr(), 23, 2)
    inherited = document.add_paragraph("原生样式编号条款", style="ContractNativeList")
    document.save(path)

    parsed = NativeContractParser().parse(path, generation_id="generation-numbering")

    assert parsed.blocks[0].metadata["native_numbering"] == {
        "mode": "NATIVE",
        "effective_num_id": 17,
        "list_level": 1,
        "source": "direct",
    }
    assert parsed.blocks[1].metadata["native_numbering"] == {
        "mode": "NATIVE",
        "effective_num_id": 23,
        "list_level": 2,
        "source": "style",
    }
    assert parsed.blocks[0].metadata["container_path"] == "document/body"


def test_scanned_pdf_is_rejected_with_frozen_error(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ContractError) as captured:
        NativeContractParser().parse(path, generation_id="generation-scanned")

    assert captured.value.code == "SCANNED_DOCUMENT_UNSUPPORTED"


def test_mixed_pdf_requires_ocr_preprocessing(tmp_path: Path) -> None:
    reader = PdfReader(io.BytesIO(text_pdf_bytes("Article 1: Party A shall pay on time.")))
    writer = PdfWriter()
    writer.add_page(reader.pages[0])
    writer.add_blank_page(width=612, height=792)
    path = tmp_path / "mixed.pdf"
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ContractError) as captured:
        NativeContractParser().parse(path, generation_id="generation-mixed")

    assert captured.value.code == "OCR_PREPROCESS_REQUIRED"


def test_encrypted_pdf_is_rejected_with_frozen_error(tmp_path: Path) -> None:
    path = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    with path.open("wb") as handle:
        writer.write(handle)

    with pytest.raises(ContractError) as captured:
        NativeContractParser().parse(path, generation_id="generation-encrypted")

    assert captured.value.code == "FILE_ENCRYPTED"


@pytest.mark.parametrize(
    ("filename", "content", "expected_code"),
    [
        ("broken.pdf", b"%PDF-not-valid", "FILE_CORRUPTED"),
        ("broken.docx", b"PK-not-valid", "FILE_CORRUPTED"),
        ("encrypted.docx", bytes.fromhex("D0CF11E0A1B11AE1") + b"encrypted", "FILE_ENCRYPTED"),
    ],
)
def test_invalid_documents_map_to_stable_errors(
    tmp_path: Path,
    filename: str,
    content: bytes,
    expected_code: str,
) -> None:
    path = tmp_path / filename
    path.write_bytes(content)

    with pytest.raises(ContractError) as captured:
        NativeContractParser().parse(path, generation_id="generation-invalid")

    assert captured.value.code == expected_code


def test_valid_docx_package_is_not_misclassified_as_encrypted(tmp_path: Path) -> None:
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("合同正文")
    document.save(buffer)
    path = tmp_path / "valid.docx"
    path.write_bytes(buffer.getvalue())

    parsed = NativeContractParser().parse(path, generation_id="generation-valid")

    assert parsed.blocks[0].text == "合同正文"


def test_contract_parser_does_not_accept_proof_text_formats(tmp_path: Path) -> None:
    path = tmp_path / "contract.txt"
    path.write_text("合同正文", "utf-8")

    with pytest.raises(ContractError) as captured:
        NativeContractParser().parse(path, generation_id="generation-text")

    assert captured.value.code == "FILE_TYPE_UNSUPPORTED"
