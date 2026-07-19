from __future__ import annotations

import io
import os
import uuid
from pathlib import Path

import psycopg
import pytest
from docx import Document

from contract.application.document_processing import ContractDocumentProcessor
from contract.application.ports import UploadedContract
from contract.config import Settings
from contract.persistence.models import DocumentBlockCreate
from contract.persistence.postgres.migrate import run_migrations
from contract.persistence.postgres.repository import ContractRepository


DATABASE_URL = os.getenv("CONTRACT_TEST_DATABASE_URL", "")
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.mark.skipif(not DATABASE_URL, reason="CONTRACT_TEST_DATABASE_URL is not configured")
def test_document_processing_persists_parse_draft_then_activates_atomically(tmp_path: Path) -> None:
    settings = Settings(database_url=DATABASE_URL, data_dir=tmp_path)
    run_migrations(settings)
    repository = ContractRepository(settings)
    processor = ContractDocumentProcessor(settings, repository)
    marker = uuid.uuid4().hex
    tenant_id = f"tenant-{marker}"
    user_id = f"user-{marker}"
    requested_document_id = f"document-{marker}"
    requested_generation_id = f"generation-{marker}"
    upload = UploadedContract(
        filename="服务合同.docx",
        content_type=DOCX_CONTENT_TYPE,
        content=_docx_bytes(),
    )

    try:
        first = processor.process(
            upload=upload,
            document_id=requested_document_id,
            generation_id=requested_generation_id,
            tenant_id=tenant_id,
            user_id=user_id,
            contract_version_id=f"version-{marker}",
        )

        assert first.document_id == requested_document_id
        assert first.document_reused is False
        assert first.generation_reused is False
        assert first.parse_completed is False
        assert first.block_count == 3
        assert first.structural_ir is not None
        generation = repository.get_parse_generation(
            first.generation_id,
            document_id=first.document_id,
            tenant_id=tenant_id,
        )
        document = repository.get_document(
            first.document_id,
            tenant_id=tenant_id,
            user_id=user_id,
        )
        assert generation is not None and generation["status"] == "RUNNING"
        assert document is not None and document["active_generation_id"] is None
        assert (tmp_path / document["storage_path"]).read_bytes() == upload.content

        repeated = processor.process(
            upload=upload,
            document_id=f"ignored-document-{marker}",
            generation_id=f"ignored-generation-{marker}",
            tenant_id=tenant_id,
            user_id=user_id,
            contract_version_id=f"version-{marker}",
        )
        assert repeated.document_id == first.document_id
        assert repeated.generation_id == first.generation_id
        assert repeated.document_reused is True
        assert repeated.generation_reused is True

        rows = repository.list_blocks(first.generation_id, tenant_id=tenant_id)
        blocks = [
            DocumentBlockCreate(
                block_id=row["block_id"],
                block_no=row["block_no"],
                block_type=row["block_type"],
                text=row["text"],
                page_number=row["page_number"],
                paragraph_no=row["paragraph_no"],
                char_start=row["char_start"],
                char_end=row["char_end"],
                heading_path=row["heading_path"],
                metadata=row["metadata_json"],
            )
            for row in rows
        ]
        completed, reused = repository.complete_parse_generation(
            generation_id=first.generation_id,
            document_id=first.document_id,
            tenant_id=tenant_id,
            blocks=blocks,
            contract_ir=first.structural_ir,
        )
        assert reused is False
        assert completed["status"] == "SUCCEEDED"
        assert repository.get_active_generation(first.document_id, tenant_id=tenant_id)["id"] == first.generation_id

        completed_retry = processor.process(
            upload=upload,
            document_id=f"another-ignored-document-{marker}",
            generation_id=f"another-ignored-generation-{marker}",
            tenant_id=tenant_id,
            user_id=user_id,
            contract_version_id=f"version-{marker}",
        )
        assert completed_retry.generation_id == first.generation_id
        assert completed_retry.generation_reused is True
        assert completed_retry.parse_completed is True
    finally:
        with psycopg.connect(DATABASE_URL) as conn:
            conn.execute("DELETE FROM contract_document WHERE tenant_id = %s", (tenant_id,))
            conn.commit()


def _docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("第一章 服务", level=1)
    document.add_paragraph("第一条 乙方提供技术服务。")
    document.add_paragraph("第二条 甲方于验收后付款。")
    document.save(buffer)
    return buffer.getvalue()
