from __future__ import annotations

import tempfile
import uuid
from pathlib import Path
from typing import Iterable

from proof.domain import DocumentBlock, ParsedDocument
from proof.domain.numbering import (
    ARTICLE_START_RE,
    CHAPTER_RE,
    PAGE_FOOTER_RE,
    SECTION_RE,
)
from proof.errors import ProofError


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")
PARSER_VERSION = "policy-parser-v3"


def parse_document(path: Path) -> ParsedDocument:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ProofError(
            "unsupported_file_type",
            f"Unsupported file type: {extension or '<none>'}",
            status_code=415,
            details={"supported": sorted(SUPPORTED_EXTENSIONS)},
        )
    if extension == ".pdf":
        return _parse_pdf(path)
    if extension == ".docx":
        return _parse_docx(path)
    return _parse_text(path, markdown=extension == ".md")


def parse_document_bytes(content: bytes, filename: str) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ProofError(
            "unsupported_file_type",
            f"Unsupported file type: {suffix or '<none>'}",
            status_code=415,
            details={"supported": sorted(SUPPORTED_EXTENSIONS)},
        )
    with tempfile.NamedTemporaryFile(suffix=suffix) as handle:
        handle.write(content)
        handle.flush()
        return parse_document(Path(handle.name))


class _BlockBuilder:
    def __init__(self) -> None:
        self.blocks: list[DocumentBlock] = []
        self.cursor = 0
        self.paragraph_index = 0
        self.chapter = ""
        self.section = ""

    def add(
        self,
        text: str,
        *,
        page_no: int | None = None,
        block_type: str | None = None,
        metadata: dict | None = None,
    ) -> None:
        raw_value = str(text or "")
        value = raw_value.strip()
        if not value:
            return
        block_metadata = dict(metadata or {})
        block_metadata.setdefault("leading_whitespace", len(raw_value) - len(raw_value.lstrip()))
        detected_type, heading_value = _classify_block(value, block_type)
        if detected_type == "heading":
            if CHAPTER_RE.match(value):
                self.chapter = heading_value
                self.section = ""
            elif SECTION_RE.match(value):
                self.section = heading_value
        self.paragraph_index += 1
        heading_path = [item for item in (self.chapter, self.section) if item]
        start = self.cursor
        end = start + len(value)
        self.blocks.append(
            DocumentBlock(
                id=uuid.uuid4().hex,
                ordinal=len(self.blocks) + 1,
                block_type=detected_type,
                text=value,
                page_no=page_no,
                paragraph_index=self.paragraph_index,
                heading_path=heading_path,
                char_start=start,
                char_end=end,
                metadata=block_metadata,
            )
        )
        self.cursor = end + 1


def _classify_block(text: str, preferred: str | None = None) -> tuple[str, str]:
    if PAGE_FOOTER_RE.match(text.replace("\u3000", " ").strip()):
        return "footer", ""
    chapter = CHAPTER_RE.match(text)
    if chapter:
        return "heading", " ".join(part for part in chapter.groups() if part).strip()
    section = SECTION_RE.match(text)
    if section:
        return "heading", " ".join(part for part in section.groups() if part).strip()
    if preferred == "heading":
        return "heading", text.lstrip("#").strip()
    if ARTICLE_START_RE.match(text):
        return "article", ""
    return preferred or "paragraph", ""


def _parse_text(path: Path, *, markdown: bool) -> ParsedDocument:
    text = _read_text(path)
    builder = _BlockBuilder()
    for line in text.splitlines():
        value = line.rstrip()
        stripped = value.lstrip()
        preferred = "heading" if markdown and stripped.startswith("#") else None
        builder.add(stripped.lstrip("#").strip() if preferred else value, block_type=preferred)
    return ParsedDocument(builder.blocks, path.suffix.lower().lstrip("."))


def _read_text(path: Path) -> str:
    errors: list[str] = []
    for encoding in TEXT_ENCODINGS:
        try:
            return path.read_text(encoding)
        except UnicodeDecodeError as exc:
            errors.append(f"{encoding}: {exc.reason}")
    raise ProofError(
        "document_decode_failed",
        "Unable to decode the text document.",
        status_code=422,
        details={"errors": errors},
    )


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
    except Exception as exc:
        raise ProofError("document_parse_failed", f"Unable to read PDF: {exc}", status_code=422) from exc

    builder = _BlockBuilder()
    warnings: list[str] = []
    pages_with_text = 0
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            warnings.append(f"page {page_no}: {exc}")
            continue
        if page_text.strip():
            pages_with_text += 1
        for line in page_text.splitlines():
            builder.add(line, page_no=page_no)
    if pages_with_text == 0:
        raise ProofError(
            "ocr_required",
            "The PDF has no selectable text and requires OCR.",
            status_code=422,
        )
    return ParsedDocument(builder.blocks, "pdf", warnings)


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(path))
    except Exception as exc:
        raise ProofError("document_parse_failed", f"Unable to read DOCX: {exc}", status_code=422) from exc

    builder = _BlockBuilder()
    for item in _iter_docx_body(document):
        if isinstance(item, Paragraph):
            style_name = str(getattr(item.style, "name", "") or "")
            preferred = "heading" if style_name.lower().startswith("heading") or style_name.startswith("标题") else None
            paragraph_metadata = _docx_paragraph_metadata(item, style_name)
            for line in item.text.splitlines():
                builder.add(line, block_type=preferred, metadata=paragraph_metadata)
        elif isinstance(item, Table):
            rows = []
            for row in item.rows:
                values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                builder.add("\n".join(rows), block_type="table")
    return ParsedDocument(builder.blocks, "docx")


def _iter_docx_body(document) -> Iterable[object]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = document.element.body
    for child in parent.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _docx_paragraph_metadata(paragraph, style_name: str) -> dict[str, int | str | None]:
    paragraph_format = paragraph.paragraph_format
    left_indent = paragraph_format.left_indent
    first_line_indent = paragraph_format.first_line_indent
    outline_level = None
    properties = getattr(paragraph._p, "pPr", None)
    if properties is not None and properties.outlineLvl is not None:
        outline_level = int(properties.outlineLvl.val)
    return {
        "style": style_name,
        "left_indent_twips": int(left_indent.twips) if left_indent is not None else None,
        "first_line_indent_twips": int(first_line_indent.twips) if first_line_indent is not None else None,
        "outline_level": outline_level,
    }
