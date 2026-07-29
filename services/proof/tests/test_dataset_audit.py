from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from proof.application.dataset_audit import DatasetAuditor
from proof.application.service import ProofService
from proof.config import Settings
from proof.errors import ProofError
from proof.tenant import tenant_scope


def test_dataset_audit_classifies_files_and_exposes_complete_chunks(tmp_path: Path) -> None:
    directory = tmp_path / "output/01_修改后公司规则文档"
    directory.mkdir(parents=True)
    ready_path = directory / "08_公司_[融资管理制度]_制度审校测试修改版.docx"
    blocked_path = directory / "无条款.docx"
    _write_docx(ready_path, ["第一条 第一段。", "第一条 重复编号。"])
    _write_docx(blocked_path, ["只有前言，没有正式条款。"])

    auditor = DatasetAuditor(tmp_path, max_upload_bytes=1024 * 1024)
    snapshot = auditor.audit()

    assert snapshot["summary"] == {
        "candidate_file_count": 2,
        "ready_file_count": 1,
        "blocked_file_count": 1,
        "anomaly_file_count": 2,
        "block_count": 3,
        "clause_count": 2,
        "mixed_structure_count": 0,
    }
    ready = next(item for item in snapshot["files"] if item["scan_status"] == "ready")
    assert ready["title"] == "融资管理制度"
    assert ready["level_code"] == "peer"
    assert ready["category_code"] == "financing_guarantee"
    assert ready["clause_count"] == 2
    assert ready["structure_profile"] == "article"
    assert {item["code"] for item in ready["anomalies"]} == {"duplicate_clause_number", "short_clause"}

    detail = auditor.file_detail(ready["id"])
    assert [item["text"] for item in detail["clauses"]] == [
        "第一条 第一段。",
        "第一条 重复编号。",
    ]

    with pytest.raises(ProofError) as exc_info:
        auditor.file_detail("unknown")
    assert exc_info.value.code == "dataset_file_not_found"


def test_subsidiary_filename_does_not_override_title_classification(tmp_path: Path) -> None:
    directory = tmp_path / "output/03_三级子公司局部制度"
    directory.mkdir(parents=True)
    path = directory / "L3_采购管理办法_三级子公司局部制度_测试修改版.docx"
    _write_docx(path, ["第一条 采购事项由采购部门审核。"])

    item = DatasetAuditor(tmp_path, max_upload_bytes=1024 * 1024).audit()["files"][0]

    assert item["title"] == "采购管理办法"
    assert item["category_code"] == "procurement_supply"


def test_dataset_roots_are_derived_only_from_current_java_tenant(tmp_path: Path) -> None:
    service = object.__new__(ProofService)
    service.settings = Settings(_env_file=None)
    service.dataset_root = tmp_path / "datasets"
    service.dataset_auditor = None
    service._tenant_dataset_auditors = {}

    with tenant_scope("1"):
        main = service._dataset_auditor()
    with tenant_scope("2"):
        demo = service._dataset_auditor()

    assert main.root.parent == demo.root.parent == (tmp_path / "datasets" / "tenants")
    assert main.root != demo.root


def _write_docx(path: Path, lines: list[str]) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        paragraph.add_run(line)
    document.save(path)
