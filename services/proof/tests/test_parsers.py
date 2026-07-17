from __future__ import annotations

from pathlib import Path

from docx import Document

from proof.application.chunking import split_into_clause_units
from proof.infrastructure.parsers import parse_document


def test_markdown_chapters_are_metadata_not_units(tmp_path: Path) -> None:
    path = tmp_path / "policy.md"
    path.write_text(
        "# 示例制度\n\n第一章 总则\n\n第一条 第一款内容。\n补充段落。\n\n第二章 职责\n第二条 第二款内容。\n",
        "utf-8",
    )

    parsed = parse_document(path)
    units = split_into_clause_units(parsed.blocks)

    assert len(units) == 2
    assert units[0].heading_path == ["第一章 总则"]
    assert units[1].heading_path == ["第二章 职责"]
    assert "第二章 职责" not in units[0].text


def test_plain_text_accepts_gb18030(tmp_path: Path) -> None:
    path = tmp_path / "policy.txt"
    path.write_bytes("第一条 中文制度。".encode("gb18030"))
    parsed = parse_document(path)
    units = split_into_clause_units(parsed.blocks)
    assert units[0].text == "第一条 中文制度。"


def test_plain_text_keeps_leading_indentation_as_metadata(tmp_path: Path) -> None:
    path = tmp_path / "indented.txt"
    path.write_text("  1. 一级标题\n    （1）子项", "utf-8")
    parsed = parse_document(path)
    assert parsed.blocks[0].metadata["leading_whitespace"] == 2
    assert parsed.blocks[1].metadata["leading_whitespace"] == 4


def test_docx_manual_line_breaks_create_separate_article_blocks(tmp_path: Path) -> None:
    path = tmp_path / "policy.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("第一条 第一条内容。")
    paragraph.add_run().add_break()
    paragraph.add_run("第二条 第二条内容。")
    document.save(path)

    parsed = parse_document(path)
    units = split_into_clause_units(parsed.blocks)

    assert [unit.clause_no_raw for unit in units] == ["第一条", "第二条"]
    assert [unit.text for unit in units] == ["第一条 第一条内容。", "第二条 第二条内容。"]
